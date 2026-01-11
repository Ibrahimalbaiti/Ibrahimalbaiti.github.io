import logging
from typing import Optional

from arabic_reshaper import reshape
from bidi.algorithm import get_display
from PIL import ImageFont
from pptx.util import Pt

logger = logging.getLogger(__name__)


class LayoutEngine:
    def fix_rtl_text(self, text: str) -> str:
        if not text:
            return text
        reshaped = reshape(text)
        return get_display(reshaped)

    def fit_text_to_shape(self, shape, translated_text: str, font_path: Optional[str] = None) -> None:
        if not hasattr(shape, "text_frame"):
            return
        text_frame = shape.text_frame
        text_frame.text = translated_text
        width_pt = shape.width / 12700
        height_pt = shape.height / 12700

        font_size = 18
        while font_size >= 8:
            try:
                font = ImageFont.truetype(font_path or "DejaVuSans.ttf", font_size)
            except OSError:
                font = ImageFont.load_default()
            text_width, text_height = font.getsize_multiline(translated_text)
            if text_width <= width_pt and text_height <= height_pt:
                break
            font_size -= 1

        for paragraph in text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(font_size)

        if font_size < 8:
            logger.warning("Text overflow detected; truncating content")
            text_frame.text = translated_text[:200]
