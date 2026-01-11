import gc
import logging
from pathlib import Path

import torch
from celery import Celery
from pptx import Presentation

from app.core.config import settings
from app.core.converter import convert_ppt_to_pptx
from app.engines.gemini import GeminiEngine
from app.engines.layout import LayoutEngine
from app.engines.nllb import NLLBEngine
from app.engines.vision import VisionEngine

logger = logging.getLogger(__name__)
logging.basicConfig(level=logging.INFO)

celery_app = Celery("pptx", broker=settings.redis_url, backend=settings.redis_url)

OUTPUT_DIR = Path("/app/data/results")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def download_models() -> None:
    try:
        VisionEngine()
    except Exception as exc:  # noqa: BLE001
        logger.warning("Vision models not fully initialized: %s", exc)
    try:
        NLLBEngine()
    except Exception as exc:  # noqa: BLE001
        logger.warning("NLLB model not fully initialized: %s", exc)


download_models()


def _select_engine(engine_type: str):
    if engine_type.lower() == "nllb":
        return NLLBEngine()
    return GeminiEngine(settings.gemini_api_key)


def _is_arabic(lang_code: str) -> bool:
    return lang_code.lower().startswith("ar") or "arab" in lang_code.lower()


@celery_app.task(bind=True, name="process_presentation_task")
def process_presentation_task(
    self,
    file_path: str,
    source_lang: str,
    target_lang: str,
    engine_type: str,
):
    engine = _select_engine(engine_type)
    layout_engine = LayoutEngine()
    vision_engine = VisionEngine()

    try:
        path = Path(file_path)
        if path.suffix.lower() == ".ppt":
            path = Path(convert_ppt_to_pptx(str(path)))

        presentation = Presentation(str(path))
        slide_elements = []
        total_slides = len(presentation.slides)

        for index, slide in enumerate(presentation.slides, start=1):
            self.update_state(
                state="PROGRESS",
                meta={"progress": f"Processing Slide {index}/{total_slides}"},
            )
            for shape in slide.shapes:
                if shape.shape_type == 13 and hasattr(shape, "image"):
                    image_text = vision_engine.extract_text_from_image(shape.image.blob)
                    if image_text.strip():
                        translated = engine.translate_text(image_text, source_lang, target_lang)
                        slide_elements.append(
                            {
                                "slide": index,
                                "text": translated,
                                "x": float(shape.left),
                                "y": float(shape.top),
                                "w": float(shape.width),
                                "h": float(shape.height),
                            }
                        )
                    continue

                if not shape.has_text_frame:
                    continue

                original_text = shape.text_frame.text
                translated_text = engine.translate_text(original_text, source_lang, target_lang)
                if _is_arabic(target_lang):
                    translated_text = layout_engine.fix_rtl_text(translated_text)
                layout_engine.fit_text_to_shape(shape, translated_text, None)

                slide_elements.append(
                    {
                        "slide": index,
                        "text": translated_text,
                        "x": float(shape.left),
                        "y": float(shape.top),
                        "w": float(shape.width),
                        "h": float(shape.height),
                    }
                )

        output_path = OUTPUT_DIR / f"{path.stem}_translated.pptx"
        presentation.save(str(output_path))

        return {
            "output_path": str(output_path),
            "elements": slide_elements,
        }
    finally:
        if torch.cuda.is_available():
            torch.cuda.empty_cache()
        gc.collect()
